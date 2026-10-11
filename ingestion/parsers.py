"""Parsers: turn a file into a list of structural Blocks (+ a 'naive' raw text for comparison).

Every parser returns a ParseResult and NEVER raises: problems become warnings, so one bad file
cannot crash a batch ingestion job.
"""
from __future__ import annotations

import re
import statistics
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from pathlib import Path

from . import ocr as ocrmod
from .models import Block, ParseResult, warn
from .utils import LIST_RE, decode_bytes, table_to_text

IMAGE_EXT = {"png", "jpg", "jpeg", "tif", "tiff", "bmp", "webp", "gif"}
PAGE_NUM_RE = re.compile(r"^(page\s*)?\d+(\s*(of|/)\s*\d+)?$", re.I)
FOOTNOTE_START_RE = re.compile(r"^(\d{1,2}|[*†‡§])[\s.)]\s*\S")
CID_RE = re.compile(r"\(cid:(\d+)\)")
BULLET_CIDS = {"127", "149", "183"}      # glyph ids that are almost always a bullet point
SENTENCE_END_RE = re.compile(r"""[.!?:;।"')\]]$""")


# ---------------------------------------------------------------- file type detection
def detect_type(path: Path):
    """Decide the file type from its first bytes (magic numbers), not only the extension."""
    ext = path.suffix.lower().lstrip(".")
    with open(path, "rb") as f:
        head = f.read(8)
    ext_type = ("pdf" if ext == "pdf" else "docx" if ext == "docx" else "html" if ext in ("html", "htm")
                else "text" if ext in ("txt", "md") else "image" if ext in IMAGE_EXT else "doc" if ext == "doc" else None)
    sniff = None
    if head.startswith(b"%PDF"):
        sniff = "pdf"
    elif head.startswith(b"PK"):
        try:
            with zipfile.ZipFile(path) as z:
                sniff = "docx" if "word/document.xml" in z.namelist() else "zip"
        except zipfile.BadZipFile:
            sniff = None
    elif head.startswith(b"\x89PNG") or head[:3] == b"\xff\xd8\xff" or head[:4] in (b"GIF8", b"II*\x00", b"MM\x00*", b"RIFF"):
        sniff = "image"
    elif head.startswith(b"\xd0\xcf\x11\xe0"):
        sniff = "doc"
    warnings = []
    final = sniff or ext_type
    if sniff and ext_type and sniff != ext_type:
        warnings.append(warn("warning", "EXTENSION_MISMATCH",
                             f"File extension says '.{ext}' but the content looks like '{sniff}'. Trusting the content."))
    return final, warnings


def parse_file(path, ocr_lang: str = "eng") -> ParseResult:
    path = Path(path)
    try:
        ftype, w = detect_type(path)
        parsers = {"pdf": lambda: parse_pdf(path, ocr_lang), "docx": lambda: parse_docx(path),
                   "html": lambda: parse_html(path), "text": lambda: parse_text(path),
                   "image": lambda: parse_image(path, ocr_lang)}
        if ftype == "doc":
            res = ParseResult([], "", [warn("error", "UNSUPPORTED_TYPE", "Old binary .doc files are not supported. Open it in Word and 'Save As' .docx.")], {}, "none")
        elif ftype in parsers:
            res = parsers[ftype]()
        else:
            res = ParseResult([], "", [warn("error", "UNSUPPORTED_TYPE", f"Unsupported file type: '{path.suffix}'. Supported: PDF, DOCX, HTML, TXT/MD, PNG/JPG/TIFF.")], {}, "none")
        res.warnings = w + res.warnings
        res.doc_meta["file_type"] = ftype or "unknown"
        return res
    except Exception as e:  # last-resort safety net
        return ParseResult([], "", [warn("error", "PARSE_FAILED", f"{type(e).__name__}: {e}")], {"file_type": "unknown"}, "none")


# ---------------------------------------------------------------- PDF
def _ocr_page(page, pno, lang, warnings):
    st = ocrmod.ocr_status()
    if not st["available"]:
        warnings.append(warn("error", "OCR_UNAVAILABLE", f"Page has no text layer (scanned?) and OCR is unavailable: {st['reason']}", pno))
        return None
    if lang not in st["languages"]:
        warnings.append(warn("warning", "OCR_LANG_MISSING", f"OCR language '{lang}' is not installed; falling back to 'eng'.", pno))
        lang = "eng"
    img = page.to_image(resolution=200).original
    blocks, raw, w, stats = ocrmod.ocr_image(img, lang, page=pno)
    warnings.append(warn("info", "OCR_FALLBACK", f"No text layer on page {pno}: used OCR ({stats['words']} words, mean confidence {stats['mean_confidence']}%).", pno))
    warnings.extend(w)
    return blocks, raw, stats


def _page_lines_and_tables(page):
    tables = []
    try:
        tables = page.find_tables()
    except Exception:
        pass
    body = page
    for t in tables:  # take table areas out so their cells are not read as loose text
        try:
            body = body.outside_bbox(t.bbox)
        except Exception:
            pass
    lines = []
    for ln in body.extract_text_lines(return_chars=True):
        chars = ln["chars"]
        sizes = [round(c["size"], 1) for c in chars if c["text"].strip()]
        if not sizes:
            continue
        bold = sum(1 for c in chars if "bold" in c["fontname"].lower()) / len(chars) > 0.6
        lines.append({"text": ln["text"].strip(), "size": statistics.median(sizes), "bold": bold,
                      "top": ln["top"], "bottom": ln["bottom"], "x0": ln["x0"]})
    tbls = []
    for t in tables:
        rows = [[(c or "").replace("\n", " ").strip() for c in r] for r in t.extract()]
        tbls.append({"top": t.bbox[1], "rows": rows})
    return lines, tbls


def parse_pdf(path: Path, ocr_lang: str = "eng") -> ParseResult:
    import pdfplumber
    warnings = []
    try:
        pdf = pdfplumber.open(str(path))
    except Exception as e:
        msg = f"{type(e).__name__}: {e}"
        if "assword" in msg or "ncrypt" in msg:
            return ParseResult([], "", [warn("error", "PDF_ENCRYPTED", "PDF is password-protected. Remove the password and re-upload.")], {}, "pdfplumber")
        return ParseResult([], "", [warn("error", "PDF_CORRUPT", f"PDF could not be opened (corrupt or truncated?): {msg[:160]}")], {}, "pdfplumber")

    pages = []   # one dict per page
    with pdf:
        if not pdf.pages:
            return ParseResult([], "", [warn("error", "PDF_EMPTY", "PDF has no pages.")], {}, "pdfplumber")
        meta = {k.lower(): str(v) for k, v in (pdf.metadata or {}).items() if v}
        for pno, page in enumerate(pdf.pages, 1):
            info = {"no": pno, "h": float(page.height), "lines": [], "tables": [], "ocr": None, "raw": ""}
            try:
                info["raw"] = page.extract_text() or ""
            except Exception as e:
                warnings.append(warn("warning", "PAGE_FAILED", f"Text extraction failed: {e}", pno))
            if len(info["raw"].strip()) < 20:     # (almost) no text layer -> scanned page -> OCR
                try:
                    info["ocr"] = _ocr_page(page, pno, ocr_lang, warnings)
                except Exception as e:
                    warnings.append(warn("error", "OCR_FAILED", f"OCR failed: {e}", pno))
                if info["ocr"]:
                    info["raw"] = info["ocr"][1]
            else:
                try:
                    info["lines"], info["tables"] = _page_lines_and_tables(page)
                except Exception as e:
                    warnings.append(warn("warning", "LAYOUT_FAILED", f"Layout analysis failed, using plain text: {e}", pno))
                    info["lines"] = [{"text": t, "size": 10, "bold": False, "top": i * 12, "bottom": i * 12 + 10, "x0": 0}
                                     for i, t in enumerate(info["raw"].splitlines())]
            pages.append(info)

    raw_text = "\n\n".join(f"--- page {p['no']} ---\n{p['raw']}" for p in pages)
    n_pages = len(pages)

    # ---- PDFs without a proper font map give "(cid:127)" instead of characters: repair bullets, drop the rest
    cid = {"bullets": 0, "dropped": 0}

    def fix_cid(m):
        if m.group(1) in BULLET_CIDS:
            cid["bullets"] += 1
            return "\u2022"
        cid["dropped"] += 1
        return ""
    for p in pages:
        for l in p["lines"]:
            l["text"] = CID_RE.sub(fix_cid, l["text"]).strip()
    if cid["bullets"] or cid["dropped"]:
        warnings.append(warn("warning", "PDF_UNMAPPED_GLYPHS",
                             f"PDF fonts lack a Unicode map: {cid['bullets']} bullet glyph(s) repaired, {cid['dropped']} unknown glyph(s) removed. "
                             "The naive text still shows them as (cid:NNN)."))

    # ---- document-wide statistics: body font size + which margin lines repeat on every page
    sizes = Counter()
    for p in pages:
        for l in p["lines"]:
            sizes[l["size"]] += len(l["text"])
    body = sizes.most_common(1)[0][0] if sizes else 10.0
    heading_sizes = sorted({l["size"] for p in pages for l in p["lines"] if l["size"] >= body * 1.15}, reverse=True)
    sig = lambda t: re.sub(r"\d+", "#", t.lower().strip())
    sig_pages = defaultdict(set)
    for p in pages:
        for l in p["lines"]:
            if l["top"] < 0.09 * p["h"] or l["bottom"] > 0.91 * p["h"]:
                sig_pages[sig(l["text"])].add(p["no"])
    repeated = {s for s, ps in sig_pages.items() if n_pages >= 2 and len(ps) >= max(2, 0.5 * n_pages)}

    def classify(l, h):
        in_margin = l["top"] < 0.09 * h or l["bottom"] > 0.91 * h
        if in_margin and (sig(l["text"]) in repeated or PAGE_NUM_RE.match(l["text"])):
            return "header" if l["top"] < 0.5 * h else "footer"
        if l["top"] > 0.72 * h and l["size"] < body * 0.88:
            return "footnote"
        if l["size"] >= body * 1.15:
            return "heading"
        if l["bold"] and len(l["text"]) <= 80 and not l["text"].rstrip().endswith((".", ",", ";")) and not LIST_RE.match(l["text"]):
            return "heading"
        if LIST_RE.match(l["text"]):
            return "list_item"
        return "paragraph"

    blocks, ocr_pages, ocr_conf = [], [], []
    for p in pages:
        if p["ocr"]:
            blocks.extend(p["ocr"][0])
            ocr_pages.append(p["no"])
            ocr_conf.append(p["ocr"][2]["mean_confidence"])
            continue
        items = [(l["top"], "line", l) for l in p["lines"]] + [(t["top"], "table", t) for t in p["tables"]]
        items.sort(key=lambda x: x[0])
        cur = None   # block being assembled: {"type","lines","size","bottom","page","level"}

        def flush():
            nonlocal cur
            if cur:
                text = "\n".join(cur["lines"])
                level = 0
                if cur["type"] == "heading":
                    level = heading_sizes.index(cur["size"]) + 1 if cur["size"] in heading_sizes else min(len(heading_sizes) + 1, 4)
                b = Block(cur["type"], text, page=cur["page"], level=level, meta={"font_size": cur["size"]})
                if cur["type"] in ("header", "footer"):
                    b.removed_reason = f"running {cur['type']} (repeated on {len(sig_pages.get(sig(text), []))}/{n_pages} pages)" if sig(text) in repeated else f"page number ({cur['type']})"
                blocks.append(b)
            cur = None

        for _, kind, obj in items:
            if kind == "table":
                flush()
                cols = max((len(r) for r in obj["rows"]), default=0)
                blocks.append(Block("table", table_to_text(obj["rows"]), page=p["no"],
                                    meta={"rows": obj["rows"], "n_rows": len(obj["rows"]), "n_cols": cols}))
                continue
            l = obj
            kind = classify(l, p["h"])
            gap = (l["top"] - cur["bottom"]) if cur else 999
            small_gap = gap <= 0.5 * l["size"]
            if kind in ("header", "footer"):
                flush()
                cur = {"type": kind, "lines": [l["text"]], "size": l["size"], "bottom": l["bottom"], "page": p["no"]}
                flush()
            elif kind == "heading":
                if cur and cur["type"] == "heading" and cur["size"] == l["size"] and small_gap:
                    cur["lines"].append(l["text"]); cur["bottom"] = l["bottom"]
                else:
                    flush(); cur = {"type": "heading", "lines": [l["text"]], "size": l["size"], "bottom": l["bottom"], "page": p["no"]}
            elif kind == "footnote":
                if cur and cur["type"] == "footnote" and small_gap and not FOOTNOTE_START_RE.match(l["text"]):
                    cur["lines"].append(l["text"]); cur["bottom"] = l["bottom"]
                else:
                    flush(); cur = {"type": "footnote", "lines": [l["text"]], "size": l["size"], "bottom": l["bottom"], "page": p["no"]}
            elif kind == "list_item":
                flush(); cur = {"type": "list_item", "lines": [l["text"]], "size": l["size"], "bottom": l["bottom"], "page": p["no"]}
            else:  # plain text line: continue the current paragraph / list item, or start a new paragraph
                if cur and cur["type"] in ("paragraph", "list_item") and small_gap:
                    cur["lines"].append(l["text"]); cur["bottom"] = l["bottom"]
                else:
                    flush(); cur = {"type": "paragraph", "lines": [l["text"]], "size": l["size"], "bottom": l["bottom"], "page": p["no"]}
        flush()

    # ---- paragraphs that continue across a page break
    merged = 0
    out = []
    for b in blocks:
        if b.removed_reason:
            out.append(b)
            continue
        prev = next((x for x in reversed(out) if not x.removed_reason and x.type != "footnote"), None)
        if (prev is not None and prev.type == "paragraph" and b.type == "paragraph" and b.page == (prev.page or 0) + 1
                and not SENTENCE_END_RE.search(prev.text.rstrip()) and b.text[:1].islower()):
            prev.text += "\n" + b.text
            prev.meta["merged_across_pages"] = [prev.page, b.page]
            merged += 1
            continue
        out.append(b)
    blocks = out
    if merged:
        warnings.append(warn("info", "PARAGRAPH_SPLIT", f"{merged} paragraph(s) continued across a page break and were merged."))

    doc_meta = {"title": meta.get("title", ""), "author": meta.get("author", ""), "pages": n_pages,
                "body_font_size": body, "ocr_pages": ocr_pages,
                "ocr_mean_confidence": round(sum(ocr_conf) / len(ocr_conf), 1) if ocr_conf else None}
    if not any(not b.removed_reason for b in blocks):
        warnings.append(warn("error", "EMPTY_DOCUMENT", "No text could be extracted from this PDF."))
    return ParseResult(blocks, raw_text, warnings, doc_meta, "pdfplumber" + (" + tesseract OCR" if ocr_pages else ""))


# ---------------------------------------------------------------- DOCX
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def parse_docx(path: Path) -> ParseResult:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    warnings = []
    try:
        d = docx.Document(str(path))
    except Exception as e:
        return ParseResult([], "", [warn("error", "DOCX_CORRUPT", f"Could not open DOCX: {type(e).__name__}: {e}")], {}, "python-docx")

    raw_text = "\n".join(p.text for p in d.paragraphs)   # naive: silently drops tables and footnotes!
    blocks, seen_hf = [], set()
    for si, sec in enumerate(d.sections, 1):
        for kind, part in (("header", sec.header), ("footer", sec.footer)):
            txt = "\n".join(p.text for p in part.paragraphs if p.text.strip()).strip()
            if txt and (kind, txt) not in seen_hf:
                seen_hf.add((kind, txt))
                blocks.append(Block(kind, txt, removed_reason=f"DOCX section {si} {kind} (repeats on every page)"))

    for child in d.element.body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "p":
            p = Paragraph(child, d)
            text = p.text
            if not text.strip():
                continue
            style = p.style.name if p.style is not None else ""
            m = re.match(r"Heading (\d)", style)
            if style == "Title":
                blocks.append(Block("heading", text, level=0, meta={"style": style}))
            elif m:
                blocks.append(Block("heading", text, level=int(m.group(1)), meta={"style": style}))
            elif style.startswith("List") or (p._p.pPr is not None and p._p.pPr.numPr is not None):
                blocks.append(Block("list_item", text, meta={"style": style}))
            else:
                blocks.append(Block("paragraph", text, meta={"style": style}))
        elif tag == "tbl":
            t = Table(child, d)
            rows, merged, seen = [], 0, set()
            for r in t.rows:
                row = []
                for c in r.cells:
                    if c._tc in seen:     # python-docx repeats a merged cell for every grid column AND row it spans
                        row.append("")
                        merged += 1
                    else:
                        seen.add(c._tc)
                        row.append(c.text.strip().replace("\n", " "))
                rows.append(row)
            blocks.append(Block("table", table_to_text(rows),
                                meta={"rows": rows, "n_rows": len(rows), "n_cols": max((len(r) for r in rows), default=0), "merged_cells": merged}))

    try:   # footnotes live in a separate XML part that python-docx does not expose
        with zipfile.ZipFile(path) as z:
            if "word/footnotes.xml" in z.namelist():
                root = ET.fromstring(z.read("word/footnotes.xml"))
                for fn in root.findall(f"{_W}footnote"):
                    fid = fn.get(f"{_W}id")
                    if fn.get(f"{_W}type") in ("separator", "continuationSeparator", "continuationNotice") or fid is None or int(fid) <= 0:
                        continue
                    text = "".join(t.text or "" for t in fn.iter(f"{_W}t")).strip()
                    if text:
                        blocks.append(Block("footnote", text, meta={"id": fid}))
    except Exception as e:
        warnings.append(warn("warning", "FOOTNOTES_FAILED", f"Could not read footnotes: {e}"))

    cp = d.core_properties
    doc_meta = {"title": cp.title or "", "author": cp.author or "", "created": str(cp.created or "")}
    if not any(not b.removed_reason for b in blocks):
        warnings.append(warn("error", "EMPTY_DOCUMENT", "No text found in this DOCX."))
    return ParseResult(blocks, raw_text, warnings, doc_meta, "python-docx")


# ---------------------------------------------------------------- HTML
_BOILER_WORDS = re.compile(r"cookie|banner|advert|sidebar|menu|popup|subscribe|newsletter", re.I)


def parse_html(path: Path) -> ParseResult:
    from bs4 import BeautifulSoup, NavigableString, Tag
    html, warnings = decode_bytes(Path(path).read_bytes())
    soup = BeautifulSoup(html, "html.parser")
    raw_text = soup.get_text("\n", strip=True)            # naive: includes menus, scripts, cookie banners ...
    meta = {"title": (soup.title.get_text(strip=True) if soup.title else ""),
            "language": (soup.html.get("lang", "") if soup.html else "")}
    for m in soup.find_all("meta"):
        if m.get("name") in ("author", "description"):
            meta[m["name"]] = m.get("content", "")
    dropped = " ".join(t.get_text(" ") for t in soup(["script", "style", "noscript", "template", "svg", "iframe"]))
    for t in soup(["script", "style", "noscript", "template", "svg", "iframe"]):
        t.decompose()

    blocks = []

    def text_of(tag):
        for s in tag.find_all("sup"):                       # <sup>1</sup> would glue onto the word: "text1"
            if s.get_text(strip=True).isdigit():
                s.replace_with(f"[{s.get_text(strip=True)}]")
        for br in tag.find_all("br"):
            br.replace_with("\n")
        return tag.get_text("")

    def classes(tag):
        return " ".join(tag.get("class", [])) + " " + (tag.get("id") or "")

    def is_footnotes(tag):
        c = classes(tag).lower()
        return "footnote" in c or "endnote" in c or tag.get("role") in ("doc-endnotes", "doc-footnotes")

    def is_boiler(tag):
        return tag.name in ("nav", "footer", "aside", "form") or bool(_BOILER_WORDS.search(classes(tag)))

    def visit(node):
        for child in node.children:
            if isinstance(child, NavigableString) or not isinstance(child, Tag):
                continue
            n = child.name
            if is_footnotes(child):
                items = child.find_all("li") or child.find_all("p")
                for it in items:
                    blocks.append(Block("footnote", text_of(it), meta={"id": it.get("id", "")}))
            elif is_boiler(child):
                btype = "header" if n == "nav" else "footer" if n == "footer" else "boilerplate"
                blocks.append(Block(btype, child.get_text(" ", strip=True)[:300], removed_reason=f"HTML boilerplate <{n}>"))
            elif n in ("h1", "h2", "h3", "h4", "h5", "h6"):
                blocks.append(Block("heading", text_of(child), level=int(n[1])))
            elif n in ("p", "blockquote", "dd", "dt"):
                blocks.append(Block("paragraph", text_of(child), meta={"tag": n}))
            elif n in ("ul", "ol"):
                for li in child.find_all("li", recursive=False):
                    blocks.append(Block("list_item", text_of(li)))
            elif n == "pre":
                blocks.append(Block("code", child.get_text(), meta={"tag": "pre"}))
            elif n == "table":
                rows = [[text_of(c).strip() for c in tr.find_all(["th", "td"])] for tr in child.find_all("tr")]
                cap = child.find("caption")
                blocks.append(Block("table", table_to_text(rows),
                                    meta={"rows": rows, "n_rows": len(rows), "n_cols": max((len(r) for r in rows), default=0),
                                          "caption": cap.get_text(strip=True) if cap else ""}))
            else:
                visit(child)

    visit(soup.body or soup)
    if not any(not b.removed_reason and b.text.strip() for b in blocks):
        warnings.append(warn("error", "EMPTY_DOCUMENT", "No readable text found in this HTML."))
    res = ParseResult(blocks, raw_text, warnings, meta, "BeautifulSoup (html.parser)")
    res.dropped_text = dropped
    return res


# ---------------------------------------------------------------- plain text / markdown
def parse_text(path: Path) -> ParseResult:
    text, warnings = decode_bytes(Path(path).read_bytes())
    blocks = []
    for chunk in re.split(r"\n\s*\n", text.strip()):
        chunk = chunk.strip("\n")
        if not chunk.strip():
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", chunk)
        if m and "\n" not in chunk:
            blocks.append(Block("heading", m.group(2).strip(), level=len(m.group(1))))
        elif all(LIST_RE.match(l) for l in chunk.splitlines()):
            blocks.extend(Block("list_item", l) for l in chunk.splitlines())
        else:
            blocks.append(Block("paragraph", chunk))
    if not blocks:
        warnings.append(warn("error", "EMPTY_DOCUMENT", "The text file is empty."))
    return ParseResult(blocks, text, warnings, {}, "plain text")


# ---------------------------------------------------------------- image (OCR)
def parse_image(path: Path, ocr_lang: str = "eng") -> ParseResult:
    from PIL import Image
    st = ocrmod.ocr_status()
    if not st["available"]:
        return ParseResult([], "", [warn("error", "OCR_UNAVAILABLE", st["reason"])], {}, "none")
    try:
        img = Image.open(path)
        img.load()
    except Exception as e:
        return ParseResult([], "", [warn("error", "IMAGE_CORRUPT", f"Image could not be opened: {e}")], {}, "Pillow")
    warnings = []
    if ocr_lang not in st["languages"]:
        warnings.append(warn("warning", "OCR_LANG_MISSING", f"OCR language '{ocr_lang}' is not installed; falling back to 'eng'."))
        ocr_lang = "eng"
    blocks, raw, w, stats = ocrmod.ocr_image(img, ocr_lang, page=1)
    warnings += w
    if not blocks:
        warnings.append(warn("error", "EMPTY_DOCUMENT", "OCR could not read any text from the image."))
    meta = {"width": img.width, "height": img.height, "ocr_pages": [1], "ocr_mean_confidence": stats["mean_confidence"],
            "ocr_language": ocr_lang}
    return ParseResult(blocks, raw, warnings, meta, "Pillow + tesseract OCR")
